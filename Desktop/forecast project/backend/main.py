import os, io, sqlite3
from datetime import date
from typing import Optional
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
import openpyxl
from pydantic import BaseModel

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
DB_PATH = os.path.join(DATA_DIR, "awards.db")


def get_db():
    os.makedirs(DATA_DIR, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def init_db():
    conn = get_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS awards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ndc TEXT NOT NULL,
            ndc_description TEXT,
            customer_name TEXT,
            customer_hierarchy TEXT,
            monthly_quantity REAL,
            valid_from TEXT,
            valid_to TEXT,
            award_status TEXT,
            is_primary_award INTEGER DEFAULT 0,
            date_entered TEXT
        );
        CREATE TABLE IF NOT EXISTS ndc_master (
            ndc TEXT PRIMARY KEY,
            description TEXT
        );
        CREATE TABLE IF NOT EXISTS customer_master (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_name TEXT,
            customer_hierarchy TEXT
        );
    """)
    conn.commit()
    conn.close()


init_db()


class AwardIn(BaseModel):
    ndc: str
    ndc_description: Optional[str] = None
    customer_name: Optional[str] = None
    customer_hierarchy: Optional[str] = None
    monthly_quantity: Optional[float] = None
    valid_from: Optional[str] = None
    valid_to: Optional[str] = None
    award_status: Optional[str] = None
    is_primary_award: Optional[bool] = False
    date_entered: Optional[str] = None


@app.get("/api/awards")
def list_awards():
    conn = get_db()
    rows = conn.execute("SELECT * FROM awards ORDER BY id DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/api/awards")
def create_award(award: AwardIn):
    conn = get_db()
    conn.execute("""
        INSERT INTO awards (ndc, ndc_description, customer_name, customer_hierarchy,
            monthly_quantity, valid_from, valid_to, award_status, is_primary_award, date_entered)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (award.ndc, award.ndc_description, award.customer_name, award.customer_hierarchy,
          award.monthly_quantity, award.valid_from, award.valid_to, award.award_status,
          1 if award.is_primary_award else 0, award.date_entered or str(date.today())))
    conn.commit()
    conn.close()
    return {"success": True}


@app.delete("/api/awards/{award_id}")
def delete_award(award_id: int):
    conn = get_db()
    conn.execute("DELETE FROM awards WHERE id = ?", (award_id,))
    conn.commit()
    conn.close()
    return {"success": True}


@app.post("/api/awards/upload")
async def upload_awards(file: UploadFile = File(...)):
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    def col(names):
        for name in names:
            if name in headers:
                return headers.index(name)
        return None

    idx = {
        "ndc":     col(["ndc"]),
        "desc":    col(["ndc description", "ndc_description", "description"]),
        "cust":    col(["customer name", "customer_name", "customer"]),
        "hier":    col(["customer hierarchy", "customer_hierarchy", "hierarchy"]),
        "qty":     col(["monthly quantity", "monthly_quantity", "quantity", "qty"]),
        "from":    col(["valid from", "valid_from"]),
        "to":      col(["valid to", "valid_to"]),
        "status":  col(["award status", "award_status", "status"]),
        "primary": col(["primary award", "is_primary_award", "primary"]),
        "date":    col(["date entered", "date_entered"]),
    }

    conn = get_db()
    inserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        def v(key):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else None

        ndc = v("ndc")
        if not ndc:
            continue

        conn.execute("""
            INSERT INTO awards (ndc, ndc_description, customer_name, customer_hierarchy,
                monthly_quantity, valid_from, valid_to, award_status, is_primary_award, date_entered)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (str(ndc).strip(), str(v("desc") or ""), str(v("cust") or ""), str(v("hier") or ""),
              v("qty"), str(v("from") or ""), str(v("to") or ""),
              str(v("status") or ""), 1 if v("primary") else 0,
              str(v("date") or date.today())))
        inserted += 1

    conn.commit()
    conn.close()
    return {"inserted": inserted}


@app.get("/api/ndc/{ndc_code}")
def lookup_ndc(ndc_code: str):
    conn = get_db()
    row = conn.execute("SELECT * FROM ndc_master WHERE ndc = ?", (ndc_code.strip(),)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "NDC not found")
    return dict(row)


@app.get("/api/ndc-master")
def list_ndc_master():
    conn = get_db()
    rows = conn.execute("SELECT * FROM ndc_master ORDER BY ndc").fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/api/ndc-master/upload")
async def upload_ndc_master(file: UploadFile = File(...)):
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    ndc_i = next((i for i, h in enumerate(headers) if h == "ndc" or (h.startswith("ndc") and "desc" not in h)), None)
    desc_i = next((i for i, h in enumerate(headers) if "desc" in h), None)

    if ndc_i is None:
        raise HTTPException(400, "Expected a column named 'NDC'")
    if desc_i is None:
        raise HTTPException(400, "Expected a column named 'Description' or 'NDC Description'")

    conn = get_db()
    upserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[ndc_i]:
            continue
        conn.execute("INSERT OR REPLACE INTO ndc_master (ndc, description) VALUES (?, ?)",
                     (str(row[ndc_i]).strip(), str(row[desc_i]).strip() if row[desc_i] else ""))
        upserted += 1
    conn.commit()
    conn.close()
    return {"upserted": upserted}


@app.get("/api/customer-master")
def list_customers():
    conn = get_db()
    rows = conn.execute("SELECT * FROM customer_master ORDER BY customer_name").fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/api/customer-master/upload")
async def upload_customer_master(file: UploadFile = File(...)):
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    name_i = next((i for i, h in enumerate(headers) if "name" in h or "customer" in h), None)
    hier_i = next((i for i, h in enumerate(headers) if "hier" in h), None)

    if name_i is None:
        raise HTTPException(400, "Expected a column named 'Customer Name'")

    conn = get_db()
    conn.execute("DELETE FROM customer_master")
    inserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[name_i]:
            continue
        name = str(row[name_i]).strip()
        hier = str(row[hier_i]).strip() if hier_i is not None and row[hier_i] else ""
        conn.execute("INSERT INTO customer_master (customer_name, customer_hierarchy) VALUES (?, ?)", (name, hier))
        inserted += 1
    conn.commit()
    conn.close()
    return {"inserted": inserted}


# Serve built frontend in production
frontend_dist = os.path.join(os.path.dirname(__file__), "..", "frontend", "dist")
if os.path.exists(frontend_dist):
    app.mount("/assets", StaticFiles(directory=os.path.join(frontend_dist, "assets")), name="assets")

    @app.get("/{full_path:path}")
    def serve_frontend(full_path: str):
        return FileResponse(os.path.join(frontend_dist, "index.html"))


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
